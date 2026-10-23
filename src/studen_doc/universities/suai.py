"""
ГУАП (SUAI): титульный лист, пресет оформления лабораторных работ и DOCX-вёрстка титульного листа.

Здесь один титульный лист - ``TitlePageSUAI``, для отчетов о лабораторных работах (основной вид в этой библиотеке).
Другие документы ГУАП (курсовая работа, ВКР, реферат) оформляются иначе: им нужен свой класс титульного листа и
своя вёрстка, они регистрируются так же, как этот (``register_title_page`` и ``render_title.register``).
"""

from __future__ import annotations

import datetime
from dataclasses import dataclass, field, replace
from typing import NamedTuple

from docx.enum.table import WD_ALIGN_VERTICAL
from docx.table import _Cell
from docx.text.paragraph import Paragraph

from studen_doc.model.report import TitlePage, register_title_page
from studen_doc.ooxml.elements import Container
from studen_doc.ooxml.fonts import apply_font
from studen_doc.ooxml.paragraph import COMPACT, COMPACT_TIGHT, anchor_to_page_bottom, set_format, set_mark_font_size
from studen_doc.ooxml.table_format import Border, add_table, set_cell_borders
from studen_doc.render.docx import Context, render_title
from studen_doc.style import GOST_732, HeadingStyle, ReportStyle, SectionRule, register_style
from studen_doc.types import Align

#: Оформление лабораторных работ ГУАП: заголовки слева с абзацным отступом, «Ход работы» нумеруется. Интервал после
#: заголовка 6 пт, у заголовков 2 уровня еще и 6 пт перед ним (как в образце кафедры).
SUAI: ReportStyle = register_style(
    "suai",
    GOST_732.evolve(
        headings={
            "structural": HeadingStyle(align="left", caps=False, first_indent_cm=1.25, after_pt=6),
            "level1": HeadingStyle(after_pt=6),
            "level2": HeadingStyle(before_pt=6, after_pt=6),
            "rules": {
                "Цель работы": SectionRule(),
                "Ход работы": SectionRule(numbered=True, structural=False),
                "Контрольные вопросы": SectionRule(),
                "Вывод": SectionRule(),
            },
        }
    ),
)


@register_title_page("suai")
@dataclass(kw_only=True)
class TitlePageSUAI(TitlePage):
    """
    Титульный лист отчёта о лабораторной работе ГУАП (не единственный возможный вид титульного листа: для курсовой
    или ВКР нужен другой класс). Надписи, размеры и вертикальные интервалы можно переопределить.

    Оформление титульного листа не зависит от стиля остального документа (``ReportStyle``). Интервалы между строками
    шапки, подписей и названия сверены с титульным листом отчёта, оформленного по образцу кафедры.
    """

    student_name: str  # "Р. Е. Волчков"
    group: str  # "4318"
    subject: str  # "ИНТЕЛЕКТУАЛЬНЫЕ ИНФОРМАЦИОННЫЕ СИСТЕМЫ"
    teacher_name: str  # "Иванов М. М."
    teacher_position: str  # "доцент, канд. техн. наук"
    work_title: str = ""  # "ЗАДАЧА РЕГРЕССИИ В ORANGE"
    work_number: int | str | None = None
    work_kind: str = "ОТЧЁТ О ЛАБОРАТОРНОЙ РАБОТЕ"
    department: str = "82"
    university: str = "ГУАП"
    protected_line: str = "ОТЧЕТ ЗАЩИЩЕН С ОЦЕНКОЙ"
    teacher_label: str = "ПРЕПОДАВАТЕЛЬ"
    student_label: str = "РАБОТУ ВЫПОЛНИЛ"
    group_label: str = "СТУДЕНТ гр. №"
    course_prefix: str = "По курсу: "
    city: str = "Санкт-Петербург"
    year: int = field(default_factory=lambda: datetime.date.today().year)
    title_bold: bool = False

    # шрифты (пт)
    text_size_pt: float = 12
    title_size_pt: float = 14
    caption_size_pt: float = 10  # подписи под линиями ("подпись, дата")

    # вертикальные интервалы (пт)
    line_gap_pt: float = 10  # после коротких строк шапки ("ОТЧЕТ ЗАЩИЩЕН С ОЦЕНКОЙ") и перед блоками
    gap_after_university_pt: float = 24  # между "ГУАП" и "КАФЕДРА"
    gap_after_header_pt: float = 60  # между "КАФЕДРА" и "ОТЧЕТ ЗАЩИЩЕН С ОЦЕНКОЙ"
    gap_before_signature_pt: float = 14  # между надписью ("ПРЕПОДАВАТЕЛЬ") и строкой с линиями
    gap_before_title_pt: float = 55  # между подписями преподавателя и названием работы
    title_gap_pt: float = 15  # между "ОТЧЁТ О ЛАБОРАТОРНОЙ РАБОТЕ" и названием
    title_line_spacing: float = 1.5  # межстрочный интервал в названии работы (оно может занять две строки)
    gap_before_course_pt: float = 42  # между названием работы и строкой "По курсу"
    gap_before_student_pt: float = 148
    caption_line_spacing: float = 0.75  # подписи под линиями жмутся к линии, как в образце
    gap_before_city_pt: float = 90  # используется, если city_at_page_bottom=False
    city_at_page_bottom: bool = True  # город и год прижаты к нижнему полю независимо от потока

    # таблицы подписей
    column_gap_cm: float = 0.4
    teacher_columns: tuple[float, float, float] = (0.40, 0.27, 0.33)  # доли ширины без промежутков
    student_columns_cm: tuple[float, float, float] = (3.7, 2.0, 4.0)  # группа, ее номер, подпись
    underline_sz: int = 4  # толщина линии под значением (1/8 пт)
    underline_color: str = "000000"


# ------------------------------------------------------------------------------- вёрстка в DOCX
class _SignCell(NamedTuple):
    text: str = ""
    underline: bool = False
    align: Align = "center"


def _line(
    doc: Container,
    text: str,
    *,
    align: Align,
    size_pt: float,
    before_pt: float = 0,
    after_pt: float = 0,
    bold: bool = False,
    line_spacing: float | None = None,
) -> Paragraph:
    par = doc.add_paragraph()
    set_format(
        par,
        base=COMPACT if line_spacing is None else replace(COMPACT, line_spacing=line_spacing),
        align=align,
        before_pt=before_pt,
        after_pt=after_pt,
    )
    if text:
        apply_font(par.add_run(text), size_pt=size_pt, bold=True if bold else None)
    return par


def _cell(
    cell: _Cell, spec: _SignCell, *, size_pt: float, underline: Border, line_spacing: float | None = None
) -> None:
    par = cell.paragraphs[0]
    base = COMPACT_TIGHT if line_spacing is None else replace(COMPACT_TIGHT, line_spacing=line_spacing)
    set_format(par, base=base, align=spec.align)
    set_mark_font_size(par, size_pt)  # пустая ячейка не должна быть выше соседней с текстом
    if spec.text:
        apply_font(par.add_run(spec.text), size_pt=size_pt)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
    if spec.underline:
        set_cell_borders(cell, bottom=underline)


def _signature_table(
    doc: Container,
    widths_cm: list[float],
    values: list[_SignCell],
    captions: list[str],
    title: TitlePageSUAI,
) -> None:
    """
    Строка значений (с линией снизу у помеченных ячеек) и строка подписей мелким шрифтом.
    """
    underline: Border = ("single", title.underline_sz, title.underline_color)
    table = add_table(
        doc,
        rows=2,
        columns=len(widths_cm),
        widths_cm=widths_cm,
        align="left",
    )

    for column, spec in enumerate(values):
        _cell(
            table.cell(0, column),
            spec,
            size_pt=title.text_size_pt,
            underline=underline,
        )

    for column, caption in enumerate(captions):
        _cell(
            table.cell(1, column),
            _SignCell(caption),
            size_pt=title.caption_size_pt,
            underline=underline,
            line_spacing=title.caption_line_spacing,
        )


@render_title.register
def _render_suai_title(title: TitlePageSUAI, ctx: Context) -> None:
    doc, width = ctx.doc, ctx.text_width_cm
    text, big, gap = title.text_size_pt, title.title_size_pt, title.line_gap_pt

    _line(
        doc,
        title.university,
        align="center",
        size_pt=text,
        after_pt=title.gap_after_university_pt,
    )
    _line(
        doc,
        f"КАФЕДРА № {title.department}",
        align="center",
        size_pt=text,
        after_pt=title.gap_after_header_pt,
    )

    _line(
        doc,
        title.protected_line,
        align="left",
        size_pt=text,
        after_pt=gap,
    )
    _line(
        doc,
        title.teacher_label,
        align="left",
        size_pt=text,
        before_pt=gap,
        after_pt=title.gap_before_signature_pt,
    )
    free = width - 2 * title.column_gap_cm
    first, second, third = (free * ratio for ratio in title.teacher_columns)
    _signature_table(
        doc,
        [first, title.column_gap_cm, second, title.column_gap_cm, third],
        [
            _SignCell(title.teacher_position, True),
            _SignCell(),
            _SignCell(underline=True),
            _SignCell(),
            _SignCell(title.teacher_name, True),
        ],
        ["должность, уч. степень, звание", "", "подпись, дата", "", "инициалы, фамилия"],
        title,
    )

    number = f" №{title.work_number}" if title.work_number not in (None, "") else ""
    _line(
        doc,
        f"{title.work_kind}{number}",
        align="center",
        size_pt=big,
        before_pt=title.gap_before_title_pt,
        after_pt=title.title_gap_pt,
        bold=title.title_bold,
    )
    if title.work_title:
        _line(
            doc,
            title.work_title,
            align="center",
            size_pt=big,
            after_pt=title.title_gap_pt,
            bold=title.title_bold,
            line_spacing=title.title_line_spacing,
        )
    _line(
        doc,
        f"{title.course_prefix}{title.subject}",
        align="center",
        size_pt=big,
        before_pt=title.gap_before_course_pt,
    )

    _line(
        doc,
        title.student_label,
        align="left",
        size_pt=text,
        before_pt=title.gap_before_student_pt,
        after_pt=title.gap_before_signature_pt,
    )
    group_cm, number_cm, signature_cm = title.student_columns_cm
    small = title.column_gap_cm
    name_cm = width - (group_cm + number_cm + signature_cm + 2 * small)
    _signature_table(
        doc,
        [group_cm, number_cm, small, signature_cm, small, name_cm],
        [
            _SignCell(title.group_label, align="left"),
            _SignCell(title.group, True),
            _SignCell(),
            _SignCell(underline=True),
            _SignCell(),
            _SignCell(title.student_name, True),
        ],
        ["", "", "", "подпись, дата", "", "инициалы, фамилия"],
        title,
    )

    city = _line(
        doc,
        f"{title.city} {title.year}",
        align="center",
        size_pt=text,
        before_pt=0 if title.city_at_page_bottom else title.gap_before_city_pt,
    )
    if title.city_at_page_bottom:
        anchor_to_page_bottom(city)
