"""
Таблицы: ширины, границы, поля ячеек, заливка, повтор заголовка.
"""

from __future__ import annotations

from collections.abc import Sequence

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.table import Table, _Cell, _Row
from lxml import etree

from studen_doc.ooxml.elements import Container, cell_xml, element, row_xml, table_xml
from studen_doc.ooxml.schema import BORDER_SIDES_ORDER, TBL_PR_ORDER, TC_PR_ORDER, ensure_child, insert_ordered
from studen_doc.types import BoxAlign

TABLE_ALIGN: dict[BoxAlign, WD_TABLE_ALIGNMENT] = {
    "left": WD_TABLE_ALIGNMENT.LEFT,
    "center": WD_TABLE_ALIGNMENT.CENTER,
    "right": WD_TABLE_ALIGNMENT.RIGHT,
}

#: Граница: (вид, толщина в 1/8 пт, цвет).
Border = tuple[str, int, str]

_TWIPS_PER_CM = 1440 / 2.54


def add_table(
    container: Container,
    *,
    rows: int,
    columns: int,
    widths_cm: Sequence[float],
    align: BoxAlign,
    style: str | None = None,
) -> Table:
    """
    Создает таблицу с заданными выравниванием и ширинами столбцов. ``container`` - документ или ячейка.
    """
    table = container.add_table(rows=rows, cols=columns)
    if style is not None:
        table.style = style
    table.alignment = TABLE_ALIGN[align]
    set_table_widths(table, widths_cm)
    return table


def table_cells(table: Table) -> list[list[_Cell]]:
    """
    Ячейки таблицы по строкам, за один проход. ``Table.cell()`` и ``Row.cells`` в python-docx каждый раз
    пересобирают сетку всей таблицы: заполнение по одной ячейке растет квадратично от числа строк.
    """
    return [[_Cell(tc, table) for tc in row_xml(row).tc_lst] for row in table.rows]


def set_table_widths(table: Table, widths_cm: Sequence[float]) -> None:
    """
    Ширины столбцов: в ячейках и в сетке таблицы. Ширина ячейки ставится прямо в ``w:tcW`` (``Cell.width`` в
    python-docx медленнее на порядок, а ячеек в больших таблицах десятки тысяч).
    """
    table.autofit = False
    twips = [str(Cm(width).twips) for width in widths_cm]
    for row in table_xml(table).tr_lst:
        for tc, width_cm, width in zip(row.tc_lst, widths_cm, twips, strict=False):
            properties = tc.find(qn("w:tcPr"))
            cell_width = properties.find(qn("w:tcW")) if properties is not None else None
            if cell_width is None:  # ячейка без ширины: обычный путь
                _Cell(tc, table).width = Cm(width_cm)
            else:
                cell_width.set(qn("w:w"), width)
    grid_columns = table_xml(table).tblGrid.findall(qn("w:gridCol"))
    for column, width in zip(grid_columns, widths_cm, strict=False):
        column.set(qn("w:w"), str(int(width * _TWIPS_PER_CM)))


def set_table_borders(table: Table, *, size: int, color: str = "auto") -> None:
    """
    Одинаковые границы у всей сетки таблицы: внешние и внутренние, толщиной ``size`` (в восьмых долях пункта).
    """
    sides = ("top", "left", "bottom", "right", "insideH", "insideV")
    edges = [element(f"w:{side}", {"w:val": "single", "w:sz": size, "w:space": 0, "w:color": color}) for side in sides]
    properties = table_xml(table).tblPr
    for existing in properties.findall(qn("w:tblBorders")):
        properties.remove(existing)
    insert_ordered(properties, element("w:tblBorders", None, *edges), TBL_PR_ORDER)


def set_cell_borders(cell: _Cell, **borders: Border | None) -> None:
    """
    ``set_cell_borders(cell, bottom=("single", 4, "000000"), top=None)``; ``None`` - убрать границу.
    """
    container = ensure_child(cell_xml(cell).get_or_add_tcPr(), "w:tcBorders", TC_PR_ORDER)
    for side, border in borders.items():
        for existing in container.findall(qn(f"w:{side}")):
            container.remove(existing)
        if border is None:
            edge = element(f"w:{side}", {"w:val": "nil"})
        else:
            kind, size, color = border
            edge = element(f"w:{side}", {"w:val": kind, "w:sz": size, "w:color": color})
        insert_ordered(container, edge, BORDER_SIDES_ORDER)


def shade_cell(cell: _Cell, fill: str) -> None:
    """
    Заливка ячейки цветом ``fill`` (hex RGB).
    """
    shading = ensure_child(cell_xml(cell).get_or_add_tcPr(), "w:shd", TC_PR_ORDER)
    for name, value in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", fill)):
        shading.set(qn(name), value)


def repeat_header(row: _Row) -> None:
    """
    Строка повторяется в начале каждой страницы, на которую переходит таблица.
    """
    _row_properties(row).append(element("w:tblHeader"))


def forbid_row_split(row: _Row) -> None:
    """
    Строка не разрывается между страницами.
    """
    _row_properties(row).append(element("w:cantSplit"))


def _row_properties(row: _Row) -> etree._Element:
    """
    ``w:trPr`` строки; новый элемент встает первым (после ``w:tblPrEx``), как требует схема. Через
    ``get_or_add_trPr`` python-docx это в несколько раз медленнее, а строк в больших таблицах тысячи.
    """
    tr = row_xml(row)
    properties = tr.find(qn("w:trPr"))
    if properties is None:
        properties = element("w:trPr")
        tr.insert(1 if tr.find(qn("w:tblPrEx")) is not None else 0, properties)
    return properties


def set_cell_margins(table: Table, *, top: int, bottom: int, left: int, right: int) -> None:
    """
    Поля внутри ячеек, в двадцатых долях пункта (twips).
    """
    sides = (("top", top), ("left", left), ("bottom", bottom), ("right", right))  # порядок задан схемой
    edges = [element(f"w:{side}", {"w:w": value, "w:type": "dxa"}) for side, value in sides]
    margins = element("w:tblCellMar", None, *edges)
    insert_ordered(table_xml(table).tblPr, margins, TBL_PR_ORDER)
