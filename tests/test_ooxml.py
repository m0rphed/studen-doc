import pytest
from docx import Document
from docx.oxml.ns import qn

from studen_doc import Report, Table
from studen_doc.ooxml.elements import cell_xml, element, paragraph_xml, table_xml
from studen_doc.ooxml.fields import add_field
from studen_doc.ooxml.numbering import LevelSpec, add_abstract_num, add_list_instance, attach_to_list
from studen_doc.ooxml.paragraph import anchor_to_page_bottom, box_paragraph, set_format
from studen_doc.ooxml.schema import PPR_ORDER, SETTINGS_ORDER, TC_PR_ORDER, ensure_child, insert_ordered
from studen_doc.ooxml.settings import enable_update_fields, set_compat_mode_15
from studen_doc.ooxml.table_format import set_cell_borders, set_cell_margins, set_table_borders, shade_cell
from studen_doc.universities.suai import SUAI


def tags(node) -> list[str]:
    return [child.tag.split("}")[1] for child in node]


def in_schema_order(node, order) -> bool:
    prefixed = [f"w:{name}" for name in tags(node)]
    return prefixed == sorted(prefixed, key=order.index)


# ------------------------------------------------------------------------------ порядок элементов
def test_insert_ordered_places_the_child_before_its_first_successor():
    doc = Document()
    par = doc.add_paragraph()
    set_format(par, align="center", after_pt=6)  # jc и spacing: python-docx расставил их сам
    box_paragraph(par, border_sz=4, color="808080", space_pt=2, fill="F5F5F5")
    anchor_to_page_bottom(par)
    properties = paragraph_xml(par).pPr
    assert tags(properties) == ["framePr", "pBdr", "shd", "spacing", "jc"]
    assert in_schema_order(properties, PPR_ORDER)


def test_insert_ordered_rejects_elements_missing_from_the_order_list():
    with pytest.raises(ValueError, match="w:bogus"):
        insert_ordered(element("w:pPr"), element("w:bogus"), PPR_ORDER)


def test_ensure_child_returns_the_existing_element_without_touching_it():
    parent = element("w:tcPr")
    first = ensure_child(parent, "w:shd", TC_PR_ORDER, {"w:fill": "FF0000"})
    assert ensure_child(parent, "w:shd", TC_PR_ORDER, {"w:fill": "00FF00"}) is first
    assert first.get(qn("w:fill")) == "FF0000" and len(parent) == 1


def test_cell_properties_stay_in_schema_order_whatever_the_call_order():
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.vertical_alignment = 1  # vAlign раньше границ и заливки
    shade_cell(cell, "DDDDDD")
    set_cell_borders(cell, bottom=("single", 4, "000000"), top=("single", 8, "FF0000"))
    properties = cell_xml(cell).tcPr
    assert in_schema_order(properties, TC_PR_ORDER) and tags(properties) == ["tcW", "tcBorders", "shd", "vAlign"]
    assert tags(properties.find(qn("w:tcBorders"))) == ["top", "bottom"]


def test_cell_borders_can_be_replaced_and_removed():
    cell = Document().add_table(rows=1, cols=1).cell(0, 0)
    set_cell_borders(cell, bottom=("single", 4, "000000"))
    set_cell_borders(cell, bottom=None)
    (edge,) = cell_xml(cell).tcPr.find(qn("w:tcBorders"))
    assert dict(edge.attrib) == {qn("w:val"): "nil"}


def test_table_margins_land_before_the_look_element():
    table = Document().add_table(rows=1, cols=1)
    set_cell_margins(table, top=20, bottom=20, left=80, right=80)
    assert tags(table_xml(table).tblPr)[-2:] == ["tblCellMar", "tblLook"]
    margins = table_xml(table).tblPr.find(qn("w:tblCellMar"))
    assert tags(margins) == ["top", "left", "bottom", "right"]


def test_update_fields_goes_into_its_place_in_settings():
    doc = Document()
    enable_update_fields(doc)
    set_compat_mode_15(doc)
    names = [f"w:{name}" for name in tags(doc.settings.element) if name != "mathPr"]
    assert names.index("w:savePreviewPicture") < names.index("w:updateFields") < names.index("w:compat")
    known = [n for n in names if n in SETTINGS_ORDER]
    assert known == sorted(known, key=SETTINGS_ORDER.index)


# ------------------------------------------------------------------------------------- поля и списки
def test_field_has_the_five_parts_and_optional_run_properties():
    doc = Document()
    par = doc.add_paragraph()
    add_field(par, r"SEQ Рисунок \* ARABIC", "3", {"bold": True, "italic": False})
    runs = paragraph_xml(par).findall(qn("w:r"))
    assert [r.find(qn("w:fldChar")).get(qn("w:fldCharType")) for r in runs if r.find(qn("w:fldChar")) is not None] == [
        "begin",
        "separate",
        "end",
    ]
    assert runs[1].find(qn("w:instrText")).text == r" SEQ Рисунок \* ARABIC "
    result = runs[3]
    assert tags(result) == ["rPr", "t"] and tags(result.find(qn("w:rPr"))) == ["b"]
    assert result.find(qn("w:t")).text == "3"


def test_list_instances_get_new_ids_and_restart_numbering():
    doc = Document()
    levels = [LevelSpec("decimal", f"%{i}.", 0, 709) for i in (1, 2)]
    abstract = add_abstract_num(doc, levels)
    first = add_list_instance(doc, abstract, levels=2)
    second = add_list_instance(doc, abstract, levels=2)
    assert second == first + 1
    numbering = doc.part.numbering_part.element
    assert tags(numbering).index("abstractNum") < tags(numbering).index("num")  # все abstractNum до первого num
    instance = next(n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == str(second))
    assert len(instance.findall(qn("w:lvlOverride"))) == 2
    par = doc.add_paragraph()
    attach_to_list(par, second, 1)
    membership = paragraph_xml(par).pPr.find(qn("w:numPr"))
    assert tags(membership) == ["ilvl", "numId"] and membership.find(qn("w:numId")).get(qn("w:val")) == str(second)


# ------------------------------------------------------------------------------------ границы таблиц
def test_table_borders_follow_the_style_setting(tmp_path):
    def borders_of(size):
        report = Report(SUAI.evolve(table={"border_sz": size}))
        report.add_section("Ход работы", Table(["a"], [["1"]], caption="Т"))
        doc = Document(str(report.build_to_docx(tmp_path / f"t{size}.docx")))
        return doc.tables[0]._tbl.tblPr.find(qn("w:tblBorders"))

    thin, thick = borders_of(4), borders_of(12)
    assert tags(thin) == ["top", "left", "bottom", "right", "insideH", "insideV"]
    assert {edge.get(qn("w:sz")) for edge in thin} == {"4"} and {edge.get(qn("w:sz")) for edge in thick} == {"12"}


def test_set_table_borders_replaces_instead_of_duplicating():
    table = Document().add_table(rows=1, cols=1)
    set_table_borders(table, size=4)
    set_table_borders(table, size=8)
    borders = table_xml(table).tblPr.findall(qn("w:tblBorders"))
    assert len(borders) == 1 and borders[0][0].get(qn("w:sz")) == "8"


@pytest.mark.parametrize("name", ["PPR_ORDER", "TBL_PR_ORDER", "TC_PR_ORDER", "STYLE_ORDER", "SETTINGS_ORDER"])
def test_order_lists_have_no_duplicates_and_use_known_prefixes(name):
    from studen_doc.ooxml import schema

    order = getattr(schema, name)
    assert len(set(order)) == len(order)
    assert all(tag.split(":")[0] in {"w", "m", "sl"} for tag in order)
