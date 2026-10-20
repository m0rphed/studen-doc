"""
Список использованных источников и ссылки на них: модель, нумерация, Markdown, библиография через pandoc.
"""

import textwrap

import demo_sources
import pytest
from docx import Document

from studen_doc import (
    Bibliography,
    BibliographyError,
    Citation,
    Cite,
    MarkdownError,
    Md,
    P,
    Report,
    ResolveError,
    Source,
    Sources,
)
from studen_doc.bibliography import DEFAULT_CSL, PandocCiteproc
from studen_doc.frontends import parse_markdown_text
from studen_doc.frontends.markdown.citations import split_citations
from studen_doc.ir import RParagraph, RSources
from studen_doc.model.inline import Text
from studen_doc.pandoc import find_pandoc
from studen_doc.resolve.sources import SourceRegistry
from studen_doc.style import GOST_732

needs_pandoc = pytest.mark.skipif(find_pandoc() is None, reason="pandoc не найден")


def joined(inlines):
    return "".join(node.text for node in inlines if isinstance(node, Text))


#: Ключи, которые знает подставной оформитель.
KNOWN = {"a", "z"}


def fake_formatter(keys):
    return {key: (Text(f"запись {key}"),) for key in keys if key in KNOWN}


def resolve_nodes(report, **kwargs):
    return report.resolve(**kwargs).nodes


def entries_of(nodes):
    (sources,) = (node for node in nodes if isinstance(node, RSources))
    return [joined(entry) for entry in sources.entries]


def paragraph_text(nodes, index=0):
    return joined([node for node in nodes if isinstance(node, RParagraph)][index].inlines)


# ------------------------------------------------------------------------------ ссылки в тексте
@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("[@a]", [(Citation("a"),)]),
        ("см. [@a; @b] тут", ["см. ", (Citation("a"), Citation("b")), " тут"]),
        ("[@a, с. 5]", [(Citation("a", "с. 5"),)]),
        ("[@a, с. 5; @b]", [(Citation("a", "с. 5"), Citation("b"))]),
        ("([@кириллица-1])", ["(", (Citation("кириллица-1"),), ")"]),
        ("[@a:b.c-d_e]", [(Citation("a:b.c-d_e"),)]),
    ],
)
def test_citations_are_found_in_text(text, expected):
    assert list(split_citations(text)) == expected


@pytest.mark.parametrize(
    "text",
    ["[скобки]", "почта me@example.com", "[почта me@example.com]", "[@]", "[@ a]", "@a", "текст @a без скобок", "[]"],
)
def test_text_without_citations_is_left_alone(text):
    assert list(split_citations(text)) == [text]


@pytest.mark.parametrize("text", ["[см. @a]", "[@a;]", "[@a b]", "[@a, ]", "[@a; см. @b]", "[@a; ]"])
def test_malformed_citations_are_errors_not_plain_text(text):
    with pytest.raises(MarkdownError, match="Не удалось разобрать ссылку"):
        list(split_citations(text))


@pytest.mark.parametrize("text", ["вот [@a", "две [@a; @b и еще [@c]", "(см. [@кириллица"])
def test_unclosed_bracket_with_a_citation_is_an_error(text):
    with pytest.raises(MarkdownError, match="Не закрыта скобка"):
        list(split_citations(text))


@needs_pandoc
def test_malformed_citation_in_markdown_is_reported_with_the_fragment():
    with pytest.raises(MarkdownError, match=r"\[см\. @a\]"):
        parse_markdown_text("# Раздел\n\nТак [см. @a] сказано.\n")
    parse_markdown_text("# Раздел\n\nТак `[см. @a]` в коде допустимо.\n")  # код не разбирается


def test_cite_accepts_keys_and_citations_and_needs_at_least_one():
    assert Cite("a", Citation("b", "с. 5")).citations == (Citation("a"), Citation("b", "с. 5"))
    with pytest.raises(ValueError, match="хотя бы один"):
        Cite()


# --------------------------------------------------------------------- номера по порядку ссылок
def test_numbers_follow_the_first_citation():
    registry = SourceRegistry()
    assert registry.cite([Citation("b")]) == "[1]"
    assert registry.cite([Citation("a"), Citation("b")]) == "[2, 1]"
    assert registry.cite([Citation("c"), Citation("a")]) == "[3, 2]"


def test_three_numbers_in_a_row_become_a_range_and_locators_use_semicolons():
    registry = SourceRegistry()
    for key in "abc":
        registry.cite([Citation(key)])
    assert registry.cite([Citation("a"), Citation("b"), Citation("c")]) == "[1–3]"
    assert registry.cite([Citation("a"), Citation("b")]) == "[1, 2]"
    assert registry.cite([Citation("a", "с. 5"), Citation("b")]) == "[1, с. 5; 2]"
    assert registry.cite([Citation("a", "с. 5")]) == "[1, с. 5]"


# --------------------------------------------------------------------------- ручной список
def test_manual_sources_are_listed_in_citation_order_then_the_rest():
    report = Report(GOST_732).add_section(
        "Введение",
        P("Сначала ", Cite("b"), ", потом ", Cite("a"), "."),
        Sources([Source("a", "А"), Source("b", "Б"), "Без ключа", Source("c", Md("В"))]),
    )
    nodes = resolve_nodes(report, inline_parser=lambda sources: {s: (Text(s),) for s in sources})
    assert entries_of(nodes) == ["Б", "А", "Без ключа", "В"]
    assert paragraph_text(nodes) == "Сначала [1], потом [2]."


def test_wrong_citations_and_lists_are_reported():
    only_text = Report(GOST_732).add_section("Введение", P(Cite("a")))
    with pytest.raises(BibliographyError, match="списка нет"):
        only_text.resolve()
    unknown = Report(GOST_732).add_section("Введение", P(Cite("zz")), Sources([Source("a", "А")]))
    with pytest.raises(BibliographyError, match="Нет описания источников"):
        unknown.resolve(source_formatter=fake_formatter)
    with pytest.raises(ResolveError, match="несколько списков"):
        Report(GOST_732).add_section("Введение", Sources(["а"]), Sources(["б"])).resolve()
    with pytest.raises(ResolveError, match="дважды"):
        Report(GOST_732).add_section("Введение", Sources([Source("a", "А"), Source("a", "Б")])).resolve()


def test_sources_from_bibliography_need_a_formatter():
    report = Report(GOST_732).add_section("Введение", P(Cite("a")), Sources())
    with pytest.raises(BibliographyError, match="библиографии нет"):
        report.resolve()
    assert entries_of(resolve_nodes(report, source_formatter=fake_formatter)) == ["запись a"]


def test_nocite_and_manual_entries_mix_with_the_bibliography():
    report = Report(GOST_732).add_section(
        "Введение",
        P(Cite("m"), Cite("a")),
        Sources([Source("m", "ручная")], nocite=["z"]),
    )
    assert entries_of(resolve_nodes(report, source_formatter=fake_formatter)) == ["ручная", "запись a", "запись z"]
    missing = Report(GOST_732).add_section("Введение", Sources(nocite=["нет"]))
    with pytest.raises(BibliographyError, match="nocite"):
        missing.resolve(source_formatter=fake_formatter)


# ---------------------------------------------------------------------------------------- Markdown
@needs_pandoc
def test_markdown_citations_and_manual_sources_with_keys():
    text = textwrap.dedent(
        """\
        # Введение

        Так сказано [@one] и вот [@two, с. 7] (скобки [не ссылка]).

        # Список использованных источников

        ::: sources
        1. []{#two} Вторая запись *с курсивом*.
        2. []{#one} Первая запись.
        3. Запись без ключа.
        :::
        """
    )
    nodes = resolve_nodes(Report(GOST_732, items=list(parse_markdown_text(text).items)))
    assert paragraph_text(nodes) == "Так сказано [1] и вот [2, с. 7] (скобки [не ссылка])."
    assert entries_of(nodes) == ["Первая запись.", "Вторая запись с курсивом.", "Запись без ключа."]


@needs_pandoc
def test_sources_block_options_and_errors():
    (section,) = parse_markdown_text('# Список\n\n::: {.sources nocite="a, b"}\n:::\n').items
    assert section.content == (Sources(nocite=["a", "b"]),)
    with pytest.raises(MarkdownError, match="нумерованного списка"):
        parse_markdown_text("# Список\n\n::: sources\nпросто текст\n:::\n")


@needs_pandoc
def test_bibliography_metadata_is_read_relative_to_the_document(tmp_path):
    path = tmp_path / "r.md"
    path.write_text("---\nbibliography: [refs.bib]\ncsl: s.csl\n---\n\n# Раздел\n\nТекст.\n", encoding="utf-8")
    report = Report.from_markdown(path, style=GOST_732)
    assert report.bibliography == Bibliography([tmp_path / "refs.bib"], tmp_path / "s.csl")
    override = Report.from_markdown(path, style=GOST_732, bibliography=Bibliography("other.bib"))
    assert override.bibliography == Bibliography("other.bib")


def test_bibliography_needs_files():
    with pytest.raises(ValueError, match="хотя бы один файл"):
        Bibliography([])


# --------------------------------------------------------------------------------- pandoc citeproc
@needs_pandoc
def test_citeproc_formats_entries_by_gost_style_without_numbers(tmp_path):
    refs = tmp_path / "refs.bib"
    refs.write_text(
        "@book{zem, author={Земсков, А. И.}, title={Электронные библиотеки}, publisher={Либерея},"
        " address={М.}, year={2003}, pagetotal={351}}\n"
        "@online{web, title={Сайт}, url={https://example.org/}, urldate={2026-07-15}}\n",
        encoding="utf-8",
    )
    format_entries = PandocCiteproc(Bibliography(refs))
    entries = format_entries(["zem", "web", "missing"])
    assert list(entries) == ["zem", "web"]
    book, site = (joined(entry) for entry in entries.values())
    assert book.startswith("Земсков А. И. Электронные библиотеки. — М.: Либерея, 2003. — 351 с.")
    assert "URL: https://example.org/ (дата обращения: 15.07.2026)" in site
    assert list(format_entries(["*"])) == ["zem", "web"]
    assert DEFAULT_CSL.is_file()


@needs_pandoc
def test_missing_bibliography_file_is_reported(tmp_path):
    with pytest.raises(BibliographyError, match="не найден"):
        PandocCiteproc(Bibliography(tmp_path / "нет.bib"))(["a"])


# ------------------------------------------------------------------------------------------ демо
@needs_pandoc
def test_demo_with_bibliography_numbers_sources_by_first_citation(tmp_path):
    doc = Document(str(demo_sources.build_report().build_to_docx(tmp_path / "demo.docx")))
    text = [p.text for p in doc.paragraphs]
    assert any("доступа к научной информации [1]." in line for line in text)
    assert any("[4, с. 130]" in line and "[1, 2, 4]" in line and "[1–3]" in line for line in text)
    start = text.index("Список использованных источников", 3) + 1
    entries = doc.paragraphs[start : start + 6]
    assert entries[0].text.startswith("Земсков")
    assert entries[3].text.startswith("Леготин")
    assert entries[5].text.startswith("Z-тест")
    assert all(p._p.pPr is not None and p._p.pPr.numPr is not None for p in entries)  # настоящий список Word


@needs_pandoc
def test_manual_demo_lists_entries_in_citation_order(tmp_path):
    doc = Document(str(demo_sources.build_manual_report().build_to_docx(tmp_path / "manual.docx")))
    text = [p.text for p in doc.paragraphs]
    assert any("[2, с. 4; 3]" in line for line in text)
    start = text.index("Список использованных источников", 3) + 1
    assert text[start].startswith("Z-тест")
    assert text[start + 4].startswith("Запись без ключа")
