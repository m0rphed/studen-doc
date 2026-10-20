import json
from pathlib import Path

import demo_report
import demo_sections
import pytest
from docx import Document
from docx.oxml.ns import qn

from fingerprint import fingerprint
from helpers import M, instructions, png, texts
from studen_doc import (
    Bullets,
    Figure,
    Formula,
    Listing,
    ListItem,
    Math,
    Md,
    Numbered,
    P,
    Report,
    Subsection,
    Table,
)
from studen_doc.render.docx import DocxRenderer, render_title
from studen_doc.style import GOST_732
from studen_doc.universities.suai import SUAI, TitlePageSUAI

GOLDEN = Path(__file__).parent / "golden"


@pytest.mark.parametrize("module", [demo_report, demo_sections], ids=["demo_report", "demo_sections"])
def test_output_matches_pre_refactoring_golden(module, tmp_path):
    """Рефакторинг не должен менять документ: сравнение с отпечатком, снятым со старой реализации."""
    path = module.build_report().build_to_docx(tmp_path / "out.docx")
    expected = json.loads((GOLDEN / f"{module.__name__}.json").read_text(encoding="utf-8"))
    assert fingerprint(path) == expected


def build_everything(style=SUAI) -> Report:
    report = Report(style, title="Т", author="А")
    report.add_title_page(
        TitlePageSUAI(
            student_name="И. И. Иванов",
            group="1",
            subject="Предмет",
            teacher_name="П. П. Петров",
            teacher_position="доцент",
            work_number=1,
        )
    )
    report.add_toc()
    report.add_section(
        "Ход работы",
        P(Md("См. рисунок [](#fig:a), таблицу [](#tab:t), формулу ([](#eq:f)) и листинг [](#lst:c).")),
        Subsection(
            "Подраздел",
            Figure(png(), "Первый", label="fig:a"),
            Table(["A", "B"], [["1", "x"]], caption="Таблица", label="tab:t"),
            Formula(r"E=mc^2", label="eq:f"),
            Listing("x = 1\nprint(x)", caption="Код", label="lst:c"),
            Subsection("Глубже", P("текст")),
        ),
    )
    return report


def test_structure_fields_and_properties(tmp_path):
    doc = Document(str(build_everything().build_to_docx(tmp_path / "a.docx")))
    assert {"1 Ход работы", "1.1 Подраздел", "1.1.1 Глубже"} <= set(texts(doc))
    fields = instructions(doc)
    assert {f.split(" ")[0] for f in fields} >= {"SEQ", "REF", "TOC"}
    assert sum(f.startswith("REF ") for f in fields) == 4
    assert len(list(doc.element.body.iter(M + "oMath"))) == 1
    assert doc.core_properties.title == "Т" and doc.core_properties.author == "А"
    assert doc.sections[0].different_first_page_header_footer is True
    assert doc.settings.element.find(qn("w:updateFields")) is not None


def test_style_changes_reach_the_document(tmp_path):
    style = SUAI.evolve(page={"margin_left_mm": 25}, text={"size_pt": 12}, figure={"caption": {"dash": "–"}})
    doc = Document(str(build_everything(style).build_to_docx(tmp_path / "b.docx")))
    assert round(doc.sections[0].left_margin.mm) == 25
    assert doc.styles["Normal"].font.size.pt == 12
    assert "Рисунок 1 – Первый" in texts(doc)


def test_gost_structural_heading_is_centered_caps():
    report = Report(GOST_732)
    report.add_section("Введение", P("текст"))
    resolved = report.resolve()
    doc = DocxRenderer().build(resolved)
    head = next(p for p in doc.paragraphs if p.text == "ВВЕДЕНИЕ")
    assert head.alignment == 1  # CENTER


def test_page_number_alignment_option():
    report = Report(SUAI.evolve(page_numbers={"align": "right"}))
    report.add_section("Вывод", P("x"))
    doc = DocxRenderer().build(report.resolve())
    assert doc.sections[0].footer.paragraphs[0].alignment == 2  # RIGHT


# ------------------------------------------------------------------------------- списки Word
def list_info(par) -> tuple[int, int] | None:
    """(numId, уровень) абзаца, если он пункт списка Word."""
    num_pr = par._p.pPr.numPr if par._p.pPr is not None else None
    return None if num_pr is None else (num_pr.numId.val, num_pr.ilvl.val)


def level_of(doc, num_id: int, level: int):
    """Описание уровня списка из numbering.xml: (numFmt, lvlText, left, firstLine, suff)."""
    numbering = doc.part.numbering_part.element
    num = next(n for n in numbering.findall(qn("w:num")) if int(n.get(qn("w:numId"))) == num_id)
    abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    abstract = next(a for a in numbering.findall(qn("w:abstractNum")) if a.get(qn("w:abstractNumId")) == abstract_id)
    lvl = next(item for item in abstract.findall(qn("w:lvl")) if int(item.get(qn("w:ilvl"))) == level)
    ind = lvl.find(qn("w:pPr")).find(qn("w:ind"))
    return (
        lvl.find(qn("w:numFmt")).get(qn("w:val")),
        lvl.find(qn("w:lvlText")).get(qn("w:val")),
        int(ind.get(qn("w:left"))),
        int(ind.get(qn("w:firstLine"))),
        lvl.find(qn("w:suff")).get(qn("w:val")),
    )


def test_lists_are_real_word_lists_without_typed_markers():
    report = Report(SUAI)
    report.add_section("Вывод", Bullets(["а;", "б."]), Numbered(["раз", ListItem("два", ["2.1", "2.2"])], marker="а)"))
    doc = DocxRenderer().build(report.resolve())
    items = {p.text: list_info(p) for p in doc.paragraphs if list_info(p)}
    assert list(items) == ["а;", "б.", "раз", "два", "2.1", "2.2"]  # в тексте нет «–» и «а)» — их рисует Word
    bullet_id, letters_id = items["а;"][0], items["раз"][0]
    assert bullet_id != letters_id and items["б."][0] == bullet_id
    assert [items[t][1] for t in ("раз", "два", "2.1", "2.2")] == [0, 0, 1, 1]
    assert {items[t][0] for t in ("раз", "два", "2.1", "2.2")} == {letters_id}
    assert level_of(doc, bullet_id, 0) == ("bullet", "–", 0, 709, "space")  # маркер на абзацном отступе 1,25 см
    assert level_of(doc, letters_id, 0) == ("russianLower", "%1)", 0, 709, "space")
    assert level_of(doc, letters_id, 1) == ("decimal", "%2)", 709, 709, "space")  # внутри букв — цифры


def test_every_list_restarts_numbering_from_one():
    report = Report(SUAI)
    report.add_section("Вывод", Numbered(["a", "b"]), P("между"), Numbered(["c", "d"]))
    doc = DocxRenderer().build(report.resolve())
    ids = [list_info(p)[0] for p in doc.paragraphs if list_info(p)]
    assert ids[0] == ids[1] != ids[2] == ids[3]
    numbering = doc.part.numbering_part.element
    for num_id in (ids[0], ids[2]):
        num = next(n for n in numbering.findall(qn("w:num")) if int(n.get(qn("w:numId"))) == num_id)
        assert num.find(qn("w:lvlOverride")).find(qn("w:startOverride")).get(qn("w:val")) == "1"


def test_numbering_definitions_follow_the_style():
    style = SUAI.evolve(lists={"bullet": "•", "nested_indent_cm": 2.0}, text={"first_indent_cm": 1.0})
    report = Report(style)
    report.add_section("Вывод", Bullets(["x", ListItem("y", ["z"])]))
    doc = DocxRenderer().build(report.resolve())
    num_id = next(list_info(p)[0] for p in doc.paragraphs if list_info(p))
    assert level_of(doc, num_id, 0) == ("bullet", "•", 0, 567, "space")
    assert level_of(doc, num_id, 1) == ("bullet", "•", 1134, 567, "space")


@pytest.mark.parametrize("count, word_numbering", [(7, True), (8, False)])
def test_letter_lists_longer_than_seven_use_gost_static_letters(count, word_numbering):
    report = Report(SUAI)
    report.add_section("Вывод", Numbered([f"п{i}" for i in range(1, count + 1)], marker="а)"))
    doc = DocxRenderer().build(report.resolve())
    items = [p for p in doc.paragraphs if "п" in p.text]
    assert len(items) == count
    assert all((list_info(p) is not None) == word_numbering for p in items)
    if not word_numbering:  # Word нумерует «з» восьмым пунктом, ГОСТ — «и»
        assert [p.text.split(" ")[0] for p in items] == ["а)", "б)", "в)", "г)", "д)", "е)", "ж)", "и)"]


def test_page_numbers_can_be_disabled():
    report = Report(SUAI.evolve(page_numbers={"align": "none"}))
    report.add_section("Вывод", P("x"))
    doc = DocxRenderer().build(report.resolve())
    assert not list(doc.sections[0].footer._element.iter(qn("w:instrText")))


def test_unregistered_title_page_type_gives_clear_error():
    from studen_doc import TitlePage

    class Foreign(TitlePage):
        pass

    report = Report(SUAI, title_page=Foreign())
    with pytest.raises(TypeError, match="не зарегистрирована"):
        DocxRenderer().build(report.resolve())
    assert render_title.dispatch(TitlePageSUAI) is not render_title.dispatch(Foreign)


# ------------------------------------------------------------------ формулы: один вызов на документ
class CountingConverter:
    """Подставной конвертер формул: запоминает вызовы, возвращает минимальный oMath."""

    def __init__(self) -> None:
        self.calls: list[tuple[str, ...]] = []

    def convert(self, formulas):
        from lxml import etree

        self.calls.append(tuple(formulas))
        return {
            f: etree.fromstring(f'<m:oMath xmlns:m="{M[1:-1]}"><m:r><m:t>{f}</m:t></m:r></m:oMath>') for f in formulas
        }


def test_all_formulas_of_a_document_are_converted_in_one_call():
    report = Report(SUAI)
    report.add_section(
        "Ход работы",
        P(Md("Строчная $a^2$ и повтор $a^2$ в абзаце.")),
        Formula("E=mc^2", label="eq:e", where=[("E", Md("энергия $m$")), ("c", "скорость")]),
        Table([Md("$h$")], [[Md("$t$")]], caption=Md("Подпись $cap$")),
        Bullets([Md("пункт $p$")]),
    )
    converter = CountingConverter()
    doc = DocxRenderer(math=converter).build(report.resolve())
    assert len(converter.calls) == 1
    assert sorted(converter.calls[0]) == sorted(["E=mc^2", "E", "c", "m", "a^2", "h", "t", "cap", "p"])
    assert len(list(doc.element.body.iter(M + "oMath"))) == 10  # a^2 дважды + остальные 8 по разу


def test_converter_is_not_called_for_documents_without_formulas():
    report = Report(SUAI)
    report.add_section("Вывод", P("просто текст"))
    converter = CountingConverter()
    DocxRenderer(math=converter).build(report.resolve())
    assert converter.calls == []


def test_conversion_error_stops_the_build_before_any_drawing():
    from studen_doc import MathConversionError

    class Failing:
        def convert(self, formulas):
            raise MathConversionError("не смог")

    report = Report(SUAI)
    report.add_section("Вывод", P(Math("x")))
    with pytest.raises(MathConversionError):
        DocxRenderer(math=Failing()).build(report.resolve())


# ----------------------------------------------------------------------------------- оглавление
def _breaks(doc) -> list[int]:
    return [i for i, p in enumerate(doc.paragraphs) if p._p.xpath('.//w:br[@w:type="page"]')]


def _toc_report(**toc_options) -> Report:
    report = Report(SUAI)
    report.add_toc(**toc_options)
    report.add_section("Цель работы", P("текст"))
    return report


def test_toc_is_followed_by_a_page_break_by_default():
    doc = DocxRenderer().build(_toc_report().resolve())
    field_index = next(i for i, p in enumerate(doc.paragraphs) if p._p.xpath(".//w:instrText"))
    assert _breaks(doc) == [field_index + 1]
    assert doc.paragraphs[field_index + 2].text == "Цель работы"


def test_toc_page_break_can_be_disabled():
    doc = DocxRenderer().build(_toc_report(page_break_after=False).resolve())
    assert _breaks(doc) == []


def test_report_without_toc_has_no_toc_field_and_same_numbering(tmp_path):
    def build(with_toc: bool) -> Report:
        report = Report(SUAI)
        if with_toc:
            report.add_toc()
        report.add_section("Ход работы", Subsection("Первый", Figure(png(), "а", label="fig:a")))
        return report

    without, with_toc = build(False).resolve(), build(True).resolve()
    assert without.labels["fig:a"].display == with_toc.labels["fig:a"].display == "1"
    headings = lambda r: [(n.number, n.title) for n in r.nodes if type(n).__name__ == "RHeading"]  # noqa: E731
    assert headings(without) == headings(with_toc)
    doc = DocxRenderer().build(without)
    assert "TOC" not in " ".join(instructions(doc))
    assert doc.settings.element.find(qn("w:updateFields")) is None


# --------------------------------------------------------- старый формат списков (word_numbering=False)
STATIC_LISTS = SUAI.evolve(lists={"word_numbering": False})


def test_static_lists_output_matches_golden(tmp_path):
    """Старый формат: эталон снят и сверен с режимом Word — отличаются только абзацы-пункты списков."""
    path = demo_report.build_report(STATIC_LISTS).build_to_docx(tmp_path / "static.docx")
    expected = json.loads((GOLDEN / "demo_report_static_lists.json").read_text(encoding="utf-8"))
    assert fingerprint(path) == expected


def test_word_numbering_is_on_by_default_and_a_style_flag():
    assert SUAI.lists.word_numbering is True
    assert STATIC_LISTS.lists.word_numbering is False
    assert STATIC_LISTS.lists.bullet == SUAI.lists.bullet  # остальные настройки списков сохранены


def test_static_lists_put_markers_into_the_text_and_use_no_word_numbering():
    report = Report(STATIC_LISTS)
    report.add_section(
        "Вывод",
        Bullets(["а;", ListItem("б:", ["б1"])]),
        Numbered(["раз", ListItem("два", ["2.1"])], marker="1."),
        Numbered(["раз", ListItem("два", ["2.1"])], marker="а)"),  # до 7 пунктов, но в статическом режиме - текст
    )
    doc = DocxRenderer().build(report.resolve())
    items = [p for p in doc.paragraphs if " " in p.text]
    assert [p.text for p in items] == [
        "– а;",
        "– б:",
        "– б1",
        "1. раз",
        "2. два",
        "1. 2.1",
        "а) раз",
        "б) два",
        "1) 2.1",
    ]
    assert all(list_info(p) is None for p in items)  # numPr нет
    assert [(p.paragraph_format.first_line_indent.pt, p.paragraph_format.left_indent.pt) for p in items[:3]] == [
        (35.45, 0.0),
        (35.45, 0.0),
        (35.45, 35.45),
    ]  # отступы заданы у абзаца
    assert not list(doc.element.body.iter(qn("w:numPr")))


def test_static_mode_does_not_touch_numbering_definitions():
    from docx import Document

    baseline = len(Document().part.numbering_part.element.findall(qn("w:abstractNum")))
    report = Report(STATIC_LISTS)
    report.add_section("Вывод", Bullets(["x"]), Numbered(["y"]))
    doc = DocxRenderer().build(report.resolve())
    assert len(doc.part.numbering_part.element.findall(qn("w:abstractNum"))) == baseline


def test_static_letter_list_uses_gost_letters_for_any_length():
    report = Report(STATIC_LISTS)
    report.add_section("Вывод", Numbered([f"п{i}" for i in range(1, 12)], marker="а)"))
    doc = DocxRenderer().build(report.resolve())
    labels = [p.text.split(" ")[0] for p in doc.paragraphs if " " in p.text]
    assert labels == ["а)", "б)", "в)", "г)", "д)", "е)", "ж)", "и)", "к)", "л)", "м)"]
