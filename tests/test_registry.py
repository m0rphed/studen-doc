"""
Реестры стилей и титульных листов, конвейер сборки и подключение своего вуза без правки библиотеки.
"""

import dataclasses
import textwrap

import pytest
from docx import Document

from studen_doc import (
    GOST_732,
    MarkdownError,
    Md,
    P,
    Report,
    ResolveError,
    TitlePage,
    register_title_page,
)
from studen_doc.frontends import load_report
from studen_doc.model.report import title_page_named, title_page_names, unregister_title_page
from studen_doc.pipeline import build_docx, resolve
from studen_doc.render.docx import Context, render_title
from studen_doc.resolve import resolve_report
from studen_doc.style import register_style, style_named, style_names, unregister_style
from studen_doc.universities import SUAI, TitlePageSUAI


# ------------------------------------------------------------------------------------- стили
def test_builtin_styles_are_registered_under_their_names():
    assert style_named("gost_732") is GOST_732 and style_named("suai") is SUAI
    assert {"gost_732", "suai"} <= set(style_names())


def test_style_registration_returns_the_style_and_rejects_duplicates(temporary_style):
    custom = temporary_style("test-registry-custom", GOST_732.evolve(text={"size_pt": 13}))
    assert style_named("test-registry-custom") is custom and custom.text.size_pt == 13
    with pytest.raises(ValueError, match="уже зарегистрирован"):
        register_style("test-registry-custom", GOST_732)
    with pytest.raises(KeyError):
        style_named("нет-такого")


# --------------------------------------------------------------------------- титульные листы
def test_builtin_title_page_is_registered():
    assert title_page_named("suai") is TitlePageSUAI and "suai" in title_page_names()


def test_title_page_registration_returns_the_class_and_rejects_duplicates(temporary_title_page):
    @temporary_title_page("test-registry-page")
    @dataclasses.dataclass
    class Page(TitlePage):
        name: str

    assert title_page_named("test-registry-page") is Page
    with pytest.raises(ValueError, match="уже зарегистрирован"):

        @register_title_page("test-registry-page")
        class Other(TitlePage):
            pass

    with pytest.raises(KeyError):
        title_page_named("нет-такого")


def test_own_university_works_from_markdown_without_touching_the_library(
    tmp_path, temporary_style, temporary_title_page
):
    """
    Стиль, титульный лист и его вёрстка в одном месте: Markdown находит их по именам.
    """
    custom = temporary_style("test-registry-uni", GOST_732.evolve(text={"font_name": "Arial"}))

    @temporary_title_page("test-registry-uni")
    @dataclasses.dataclass(kw_only=True)
    class TitlePageUni(TitlePage):
        student: str
        year: int = 2000

    @render_title.register
    def _draw(title: TitlePageUni, ctx: Context) -> None:
        ctx.doc.add_paragraph(f"Титул: {title.student} {title.year}")

    path = tmp_path / "r.md"
    path.write_text(
        textwrap.dedent(
            """\
            ---
            style: test-registry-uni
            title_page:
              type: test-registry-uni
              student: Иванов
              year: 2026
            ---

            # Введение

            Текст.
            """
        ),
        encoding="utf-8",
    )
    report = load_report(path)
    assert report.style is custom and report.title_page == TitlePageUni(student="Иванов", year=2026)
    texts = [p.text for p in Document(str(report.build_to_docx(tmp_path / "out.docx"))).paragraphs]
    assert "Титул: Иванов 2026" in texts


@pytest.mark.parametrize(
    ("meta", "message"),
    [
        ("style: нет-такого\n", r"Неизвестный стиль 'нет-такого'; известны \["),
        ("style: suai\ntitle_page:\n  type: нет-такого\n", r"известны \['"),
    ],
)
def test_unknown_names_in_metadata_list_the_registered_ones(meta, message, tmp_path):
    path = tmp_path / "r.md"
    path.write_text(f"---\n{meta}---\n\n# Раздел\n", encoding="utf-8")
    with pytest.raises(MarkdownError, match=message):
        load_report(path)


# ------------------------------------------------------------------------------------- конвейер
def test_report_methods_are_thin_wrappers_over_the_pipeline(tmp_path):
    report = Report(SUAI).add_section("Введение", P("текст"))
    assert report.resolve().nodes == resolve(report).nodes
    out = report.build_to_docx(tmp_path / "a.docx")
    assert build_docx(report, tmp_path / "b.docx").read_bytes()[:2] == out.read_bytes()[:2] == b"PK"


def test_pipeline_uses_pandoc_for_md_by_default_and_numbering_alone_needs_a_parser():
    report = Report(SUAI).add_section("Введение", P(Md("*курсив*")))
    (paragraph,) = [n for n in resolve(report).nodes if type(n).__name__ == "RParagraph"]
    assert paragraph.inlines[0].italic
    with pytest.raises(ResolveError, match="inline_parser"):
        resolve_report(report)  # без конвейера никто не умеет разбирать Md



def test_unregistered_names_disappear_from_registries(temporary_style, temporary_title_page):
    temporary_style("test-registry-temp", GOST_732)

    @temporary_title_page("test-registry-temp")
    @dataclasses.dataclass
    class Page(TitlePage):
        name: str

    unregister_style("test-registry-temp")
    unregister_title_page("test-registry-temp")
    assert "test-registry-temp" not in style_names() and "test-registry-temp" not in title_page_names()
    with pytest.raises(KeyError):
        unregister_style("test-registry-temp")
    with pytest.raises(KeyError):
        unregister_title_page("test-registry-temp")
    register_style("test-registry-temp", GOST_732)  # имя снова свободно
    unregister_style("test-registry-temp")
