from docx import Document
from docx.oxml.ns import qn

from studen_doc import Formula, Listing, P, Report, Table
from studen_doc.style import GOST_732, ReportStyle, gost_732
from studen_doc.universities.suai import SUAI, TitlePageSUAI


def _doc(report: Report, tmp_path):
    return Document(str(report.build_to_docx(tmp_path / "out.docx")))


def _frames(doc):
    return [p for p in doc.paragraphs if p._p.pPr is not None and p._p.pPr.find(qn("w:framePr")) is not None]


def _title(**overrides) -> TitlePageSUAI:
    return TitlePageSUAI(
        student_name="И. И. Иванов",
        group="1",
        subject="S",
        teacher_name="T",
        teacher_position="P",
        **overrides,
    )


def test_factory_without_arguments_equals_defaults():
    assert gost_732() == ReportStyle()
    assert ReportStyle() == GOST_732


def test_factory_derives_caption_size_from_font_size():
    style = gost_732(font_size_pt=12, margins_mm=(25, 10, 15, 15))
    assert style.text.size_pt == 12
    assert style.figure.caption.size_pt == style.table.caption.size_pt == style.listing.caption.size_pt == 12
    assert (style.page.margin_left_mm, style.page.margin_bottom_mm) == (25, 15)
    assert style.evolve(text={"line_spacing": 1.0}).text.size_pt == 12


def test_table_cell_padding_comes_from_style(tmp_path):
    report = Report(SUAI.evolve(table={"cell_padding": {"left": 200}}))
    report.add_section("Раздел", Table(["a", "b"], [["1", "2"]]))
    margins = _doc(report, tmp_path).tables[0]._tbl.tblPr.find(qn("w:tblCellMar"))
    assert margins.find(qn("w:left")).get(qn("w:w")) == "200"
    assert margins.find(qn("w:top")).get(qn("w:w")) == "20"


def test_listing_border_comes_from_style(tmp_path):
    report = Report(SUAI.evolve(listing={"border_sz": 12, "border_color": "FF0000"}))
    report.add_section("Раздел", Listing("x = 1", language="python"))
    doc = _doc(report, tmp_path)
    borders = [p._p.pPr.find(qn("w:pBdr")) for p in doc.paragraphs if p._p.pPr is not None]
    top = next(b for b in borders if b is not None).find(qn("w:top"))
    assert (top.get(qn("w:sz")), top.get(qn("w:color"))) == ("12", "FF0000")


def test_formula_legend_word_and_dash_come_from_style(tmp_path):
    report = Report(SUAI.evolve(formula={"legend_word": "here", "legend_dash": "-"}))
    report.add_section("Раздел", Formula("a=b", where=[("a", "first")]))
    texts = [p.text for p in _doc(report, tmp_path).paragraphs]
    assert any(t.startswith("here ") and " - first" in t for t in texts)


def test_title_page_city_is_anchored_to_page_bottom(tmp_path):
    report = Report(SUAI)
    report.add_title_page(_title())
    report.add_section("Раздел", P("текст"))
    frames = _frames(_doc(report, tmp_path))
    assert len(frames) == 1
    assert frames[0].text.startswith("Санкт-Петербург")
    assert frames[0]._p.pPr.find(qn("w:framePr")).get(qn("w:yAlign")) == "bottom"


def test_title_page_without_bottom_anchor_uses_gap(tmp_path):
    report = Report(SUAI)
    report.add_title_page(_title(city_at_page_bottom=False))
    report.add_section("Раздел", P("текст"))
    assert _frames(_doc(report, tmp_path)) == []
