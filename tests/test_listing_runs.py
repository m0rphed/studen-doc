"""
Листинги с готовыми фрагментами (``CodeRun``): цвет, фон и начертание не зависят от pygments и источника данных.
"""

import pytest
from docx import Document
from docx.oxml.ns import qn

from helpers import instructions
from studen_doc import CodeRun, Listing, P, Report
from studen_doc.ooxml.schema import RPR_ORDER
from studen_doc.universities.suai import SUAI

RED = "cc0000"
GREEN = "00aa00"


def build(tmp_path, *blocks):
    report = Report(SUAI, title="Тест")
    report.add_section("Ход работы", *blocks)
    return Document(str(report.build_to_docx(tmp_path / "report.docx")))


def code_paragraphs(doc):
    return [p for p in doc.paragraphs if p.style.name == "Код"]


def run_properties(run):
    return run._r.find(qn("w:rPr"))


def test_styled_listing_builds_code_from_the_runs():
    listing = Listing.styled([[CodeRun("a = "), CodeRun("1", color=RED)], [], [CodeRun("end")]], "Подпись")
    assert listing.code == "a = 1\n\nend"
    assert listing.caption == "Подпись"
    assert listing.runs is not None
    assert [len(line) for line in listing.runs] == [2, 0, 1]


def test_code_that_does_not_match_the_runs_is_rejected():
    with pytest.raises(ValueError, match=r"Listing\.styled"):
        Listing("other", runs=((CodeRun("a"),),))


def test_a_run_cannot_contain_a_line_break():
    with pytest.raises(ValueError, match="перевод строки"):
        Listing.styled([[CodeRun("a\nb")]])


def test_listing_without_runs_is_unchanged():
    listing = Listing("x = 1", "Подпись", language="python")
    assert listing.runs is None


def test_each_line_is_a_paragraph_and_runs_keep_their_text(tmp_path):
    doc = build(tmp_path, Listing.styled([[CodeRun("a = "), CodeRun("1")], [CodeRun("b")]], "Вывод"))
    paragraphs = code_paragraphs(doc)
    assert [p.text for p in paragraphs] == ["a = 1", "b"]
    assert [r.text for r in paragraphs[0].runs] == ["a = ", "1"]
    assert any(p.text.startswith("Листинг 1") for p in doc.paragraphs)


def test_blank_lines_and_indentation_are_kept_as_given(tmp_path):
    lines = [[], [CodeRun("    indented")], [CodeRun("tail")], []]
    doc = build(tmp_path, Listing.styled(lines, "Вывод"))
    assert [p.text for p in code_paragraphs(doc)] == ["", "    indented", "tail", ""]


def test_color_bold_italic_underline_strike_and_background_reach_the_run(tmp_path):
    styled = CodeRun("x", color=RED, background=GREEN, bold=True, italic=True, underline=True, strike=True)
    doc = build(tmp_path, Listing.styled([[styled, CodeRun("plain")]], "Вывод"))
    first, second = code_paragraphs(doc)[0].runs
    assert str(first.font.color.rgb).lower() == RED
    assert first.font.bold and first.font.italic and first.font.underline and first.font.strike
    shading = run_properties(first).find(qn("w:shd"))
    assert shading is not None
    assert shading.get(qn("w:fill")) == GREEN
    assert shading.get(qn("w:val")) == "clear"
    assert run_properties(second) is None  # у обычного фрагмента лишних свойств нет


def test_runs_without_a_background_get_no_shading(tmp_path):
    doc = build(tmp_path, Listing.styled([[CodeRun("x", color=RED, bold=True)]], "Вывод"))
    properties = run_properties(code_paragraphs(doc)[0].runs[0])
    assert properties.find(qn("w:shd")) is None
    assert properties.find(qn("w:strike")) is None
    assert properties.find(qn("w:u")) is None


def test_run_properties_follow_the_schema_order(tmp_path):
    styled = CodeRun("x", color=RED, background=GREEN, bold=True, italic=True, underline=True, strike=True)
    doc = build(tmp_path, Listing.styled([[styled]], "Вывод"))
    names = [f"w:{child.tag.rpartition('}')[2]}" for child in run_properties(code_paragraphs(doc)[0].runs[0])]
    assert names == sorted(names, key=RPR_ORDER.index)
    assert {"w:b", "w:i", "w:strike", "w:color", "w:u", "w:shd"} <= set(names)


def test_line_numbers_work_with_runs(tmp_path):
    lines = [[CodeRun("a")], [CodeRun("b")]]
    doc = build(tmp_path, Listing.styled(lines, "Вывод", line_numbers=True))
    assert [p.text for p in code_paragraphs(doc)] == ["1  a", "2  b"]


def test_runs_replace_syntax_highlighting(tmp_path):
    listing = Listing("x = 1", "Вывод", language="python", highlight=True, runs=((CodeRun("x = 1", color=RED),),))
    doc = build(tmp_path, listing)
    (run,) = code_paragraphs(doc)[0].runs
    assert str(run.font.color.rgb).lower() == RED


def test_listing_with_runs_can_be_referenced(tmp_path):
    listing = Listing.styled([[CodeRun("a")]], "Вывод", label="lst:out")
    doc = build(tmp_path, listing, P("см. ", listing.ref))
    assert any(text.startswith("REF") for text in instructions(doc))  # ссылка стала полем с номером блока
