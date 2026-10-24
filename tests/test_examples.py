import os
import re

import pytest
from docx import Document

from helpers import M

pytest.importorskip("numpy")
pytest.importorskip("pandas")
pytest.importorskip("polars")
pytest.importorskip("scipy")
pytest.importorskip("sklearn")
pytest.importorskip("matplotlib")

import demo_api_tour
import demo_curve_fit
import demo_markdown
import demo_ml_mixed
import demo_python_loop


def build(module, tmp_path):
    return Document(str(module.build_report().build_to_docx(tmp_path / f"{module.__name__}.docx")))


def captions(doc, prefix):
    return [p.text for p in doc.paragraphs if p.text.startswith(prefix)]


def test_markdown_only_demo(tmp_path):
    doc = build(demo_markdown, tmp_path)
    assert len(captions(doc, "Таблица ")) == 1 and len(doc.inline_shapes) == 1


def test_ml_demo_mixes_markdown_with_pandas_sklearn_and_matplotlib(tmp_path):
    doc = build(demo_ml_mixed, tmp_path)
    assert len(captions(doc, "Таблица ")) == 2
    assert len(captions(doc, "Рисунок ")) == 1 and len(doc.inline_shapes) == 1
    assert len(captions(doc, "Листинг ")) == 1
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "[]{" not in text and "var=" not in text and "insert" not in text  # ничего не осталось неподставленным
    assert re.search(r"Уилка дал\s+= 0,\d{3}", text)  # p-значение подставлено с десятичной запятой
    assert list(doc.element.body.iter(M + "oMath"))


def test_python_loop_demo_generates_a_section_per_dataset(tmp_path):
    doc = build(demo_python_loop, tmp_path)
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    assert [h.split(" ", 1)[1] for h in headings if "Набор" in h] == ["Набор «Альфа»", "Набор «Бета»", "Набор «Гамма»"]
    assert len(doc.inline_shapes) == len(demo_python_loop.DATASETS)
    assert len(captions(doc, "Таблица ")) == len(demo_python_loop.DATASETS) + 1  # по набору и сводка


def test_curve_fit_demo_substitutes_numbers_and_includes_file_listing(tmp_path):
    doc = build(demo_curve_fit, tmp_path)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Выполнено 25 измерений" in text
    assert "from scipy.optimize import curve_fit" in text  # листинг подключен из файла
    assert len(doc.inline_shapes) == 1 and len(captions(doc, "Таблица ")) == 2


def test_api_tour_demo_uses_references_by_block_and_custom_mono_font(tmp_path):
    doc = build(demo_api_tour, tmp_path)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "в таблице 1, выборка показана на рисунке 1, расчет - в листинге 1." in text
    assert "вычисляются по формуле (1)" in text
    styles = doc.styles.element.xml
    assert 'w:ascii="Cascadia Code"' in styles
    assert "подготовка данных;" in text and "проверка остатков." in text  # нумерацию «а)» рисует Word
    assert len(doc.inline_shapes) == 1 and len(captions(doc, "Таблица ")) == 1


def test_terminal_demo_has_listing_and_screenshots(tmp_path):
    pytest.importorskip("termsnap")
    pytest.importorskip("PIL")
    import demo_terminal

    doc = Document(str(demo_terminal.build_report(ipython=False).build_to_docx(tmp_path / "terminal.docx")))
    assert len(captions(doc, "Листинг ")) == 1
    assert len(captions(doc, "Рисунок ")) == 3 and len(doc.inline_shapes) == 3
    code = [p for p in doc.paragraphs if p.style.name == "Код"]
    assert any(p.text.startswith("эпоха 5/5") for p in code)  # вывод остался текстом
    assert any(run.font.color.rgb is not None for p in code for run in p.runs)  # и сохранил цвета


@pytest.mark.skipif(
    os.environ.get("STUDEN_DOC_TEST_IPYTHON") != "1", reason="нужны uvx и сеть: STUDEN_DOC_TEST_IPYTHON=1"
)
def test_terminal_demo_ipython_section(tmp_path):
    pytest.importorskip("termsnap")
    import demo_terminal

    doc = Document(str(demo_terminal.build_report(ipython=True).build_to_docx(tmp_path / "terminal.docx")))
    assert len(captions(doc, "Рисунок ")) == 4 and len(captions(doc, "Листинг ")) == 2
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Out[2]: [0, 1, 4, 9, 16, 25]" in text
    assert "ZeroDivisionError: division by zero" in text
