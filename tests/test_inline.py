import re
import warnings

import pytest

from studen_doc import MarkupInPlainTextWarning, Md, P, Report, Table
from studen_doc.frontends import MarkdownError, parse_inline_markdown
from studen_doc.model.inline import B, Code, I, LineBreak, Math, Ref, Text, as_inlines, merge_text, plain_text
from studen_doc.pandoc import find_pandoc
from studen_doc.universities.suai import SUAI

needs_pandoc = pytest.mark.skipif(find_pandoc() is None, reason="pandoc не найден")


def literal(source: str):
    raise AssertionError(f"Md не ожидался: {source}")


# --------------------------------------------------------------------------- строки буквальны
@pytest.mark.parametrize(
    "source",
    ["просто текст", "**не жирный**", "[[не ссылка]]", r"\* и `не код`", "а\nб"],
)
def test_plain_strings_are_literal_text(source):
    assert as_inlines(source, literal) == (Text(source),)


# ------------------------------------------------------- формула в обычной строке: предупреждение
@pytest.mark.parametrize(
    "source",
    ["$x^2$", "матрица $A$", "в $a$ и $b$", r"$\alpha$", "$$x^2$$", "до $$\nx\n$$ после", "$a+b$, 5 $"],
)
def test_a_complete_formula_in_a_plain_string_is_warned_about_but_stays_literal(source):
    with pytest.warns(MarkupInPlainTextWarning, match="Md"):
        assert as_inlines(source, literal) == (Text(source),)


def test_the_warning_shows_the_formula_that_was_found():
    with pytest.warns(MarkupInPlainTextWarning, match=re.escape(repr("$x^2$"))):
        as_inlines("матрица $x^2$ целиком", literal)


def test_a_long_formula_is_cut_in_the_warning():
    with pytest.warns(MarkupInPlainTextWarning) as caught:
        as_inlines("$" + "x" * 100 + "$", literal)
    assert len(str(caught[0].message)) < 200 and "..." in str(caught[0].message)


@pytest.mark.parametrize(
    "source",
    ["от $5 до $10", "100$ и 200$", r"цена \$5 и \$6", "5 $ 6", "$", "$$", "$ x $", "$x ", "a $ b$", "$5$10", ""],
)
def test_dollars_that_are_not_a_formula_are_not_warned_about(source):
    with warnings.catch_warnings():
        warnings.simplefilter("error")
        as_inlines(source, literal)


def test_text_nodes_and_md_are_not_warned_about():
    with warnings.catch_warnings():
        warnings.simplefilter("error")
        as_inlines([Text("$x$"), B("$y$"), Code("$z$"), Md("$w$")], lambda source: (Math(source.strip("$")),))


def test_the_warning_can_be_switched_off():
    with warnings.catch_warnings():
        warnings.simplefilter("error")
        warnings.filterwarnings("ignore", category=MarkupInPlainTextWarning)
        as_inlines("$x^2$", literal)


def test_a_caption_and_a_cell_with_a_formula_in_a_plain_string_are_warned_about():
    report = Report(SUAI).add_section("Раздел", Table(["a"], [["$x^2$"]], caption="Матрица $A$"))
    with pytest.warns(MarkupInPlainTextWarning) as caught:
        report.resolve(inline_parser=literal)
    assert {w.message.args[0].split("'")[1] for w in caught} == {"$x^2$", "$A$"}


def test_empty_string_gives_no_nodes():
    assert as_inlines("", literal) == ()


def test_as_inlines_accepts_strings_nodes_and_mixed_sequences():
    assert as_inlines(B("x"), literal) == (Text("x", bold=True),)
    assert as_inlines(["см. ", Ref("fig:a"), I(" ещё"), " и **ещё**"], literal) == (
        Text("см. "),
        Ref("fig:a"),
        Text(" ещё", italic=True),
        Text(" и **ещё**"),
    )


def test_md_is_resolved_through_the_markup_callback():
    seen = []

    def markup(source):
        seen.append(source)
        return (Text("разобрано"),)

    assert as_inlines(["а ", Md("*б*")], markup) == (Text("а "), Text("разобрано"))
    assert seen == ["*б*"]


def test_merge_text_joins_only_same_look():
    nodes = [Text("a"), Text("b"), Text("c", bold=True), Text("d", bold=True), Code("e"), Text("f")]
    assert merge_text(nodes) == (Text("ab"), Text("cd", bold=True), Code("e"), Text("f"))
    assert merge_text([Text("x", script="sup"), Text("y")]) == (Text("x", script="sup"), Text("y"))


def test_plain_text_drops_formatting():
    assert plain_text([Text("a ", bold=True), Math("x^2"), Code(" k"), Ref("fig:a"), LineBreak()]) == "a x^2 k"


# ------------------------------------------------------------------------- разбор Md (pandoc)
@needs_pandoc
def test_md_markup_emphasis_code_math_ref():
    parsed = parse_inline_markdown(["a **b** *c* `d` $x^2$ [](#fig:z)"])
    assert parsed["a **b** *c* `d` $x^2$ [](#fig:z)"] == (
        Text("a "),
        Text("b", bold=True),
        Text(" "),
        Text("c", italic=True),
        Text(" "),
        Code("d"),
        Text(" "),
        Math("x^2"),
        Text(" "),
        Ref("fig:z"),
    )


@needs_pandoc
def test_md_ref_and_code_inherit_emphasis_and_hard_break_works():
    source = "**`код` и [](#tab:a)**\\\nнова"
    nodes = parse_inline_markdown([source])[source]
    assert nodes[:3] == (Code("код", bold=True), Text(" и ", bold=True), Ref("tab:a", bold=True))
    assert LineBreak() in nodes


@needs_pandoc
def test_md_escapes_and_lone_markers_stay_text():
    sources = [r"2 \* 3 * 4", r"цена \$5 и \$6", "5 $ 6"]
    parsed = parse_inline_markdown(sources)
    assert [plain_text(parsed[s]) for s in sources] == ["2 * 3 * 4", "цена $5 и $6", "5 $ 6"]


@needs_pandoc
def test_md_drops_spaces_and_line_breaks_at_the_edges():
    padded = "  а *б*\n "
    parsed = parse_inline_markdown([padded, "а *б*"])
    assert parsed[padded] == parsed["а *б*"]
    assert plain_text(parsed[padded]) == "а б"


@needs_pandoc
def test_a_space_between_fragments_belongs_to_a_plain_string_not_to_md():
    report = Report(SUAI).add_section("Раздел", P(Md("а"), " ", Md("б")), P(Md("а"), Md(" б")))
    paragraphs = [node for node in report.resolve().nodes if type(node).__name__ == "RParagraph"]
    assert [plain_text(paragraph.inlines) for paragraph in paragraphs] == ["а б", "аб"]


@needs_pandoc
def test_md_sources_are_parsed_by_one_process_and_duplicates_once(monkeypatch):
    from studen_doc.frontends.markdown import pandoc_ast

    calls = []
    real = pandoc_ast.run_pandoc

    def spy(args, data, **kwargs):
        calls.append(data)
        return real(args, data, **kwargs)

    monkeypatch.setattr(pandoc_ast, "run_pandoc", spy)
    parsed = parse_inline_markdown(["раз *два*", "три", "раз *два*"])
    assert len(calls) == 1 and set(parsed) == {"раз *два*", "три"}


@needs_pandoc
@pytest.mark.parametrize("source", ["", "# заголовок", "- пункт", "абзац\n\nдругой", "> цитата"])
def test_md_accepts_only_one_paragraph_of_plain_text(source):
    with pytest.raises(MarkdownError):
        parse_inline_markdown([source])


# ----------------------------------------------------------------------------- через отчет
def test_report_without_md_never_starts_the_parser():
    report = Report(SUAI)
    report.add_section("Вывод", P("**буквально**"), Table(["a"], [["1"]]))

    def boom(sources):
        raise AssertionError("парсер не нужен")

    assert report.resolve(inline_parser=boom).nodes


def test_report_collects_all_md_for_one_parser_call_and_substitutes_results():
    report = Report(SUAI)
    report.add_section("Вывод", P(Md("один"), " и ", Md("два")), Table([Md("шапка")], [[Md("ячейка")]]))
    calls = []

    def parser(sources):
        calls.append(list(sources))
        return {source: (Text(source.upper()),) for source in sources}

    resolved = report.resolve(inline_parser=parser)
    assert calls == [["один", "два", "шапка", "ячейка"]]
    paragraph = next(node for node in resolved.nodes if type(node).__name__ == "RParagraph")
    assert paragraph.inlines == (Text("ОДИН"), Text(" и "), Text("ДВА"))


def test_code_in_a_bold_table_header_is_bold(tmp_path):
    from docx import Document

    report = Report(SUAI).add_section("Раздел", Table([Md("`имя`"), "обычный"], [["x", "y"]]))
    document = Document(str(report.build_to_docx(tmp_path / "out.docx")))
    header_runs = [run for cell in document.tables[0].rows[0].cells for run in cell.paragraphs[0].runs]
    code = next(run for run in header_runs if run.text == "имя")
    assert code.bold is True
