from dataclasses import replace

import pytest

from helpers import M, png
from studen_doc import (
    B,
    Bullets,
    Code,
    Figure,
    Formula,
    I,
    LineBreak,
    Listing,
    ListItem,
    Math,
    Numbered,
    P,
    PageBreak,
    Ref,
    Report,
    Subsection,
    Table,
    Text,
)
from studen_doc.frontends import MarkdownError, parse_markdown, parse_markdown_text
from studen_doc.ir import RTable
from studen_doc.model import Section, Toc
from studen_doc.pandoc import find_pandoc
from studen_doc.style import GOST_732
from studen_doc.universities.suai import SUAI, TitlePageSUAI

pytestmark = pytest.mark.skipif(find_pandoc() is None, reason="pandoc не найден")


def parse(text: str, **kwargs):
    return parse_markdown_text(text, **kwargs)


def blocks(text: str, **kwargs):
    """
    Блоки единственного раздела.
    """
    (section,) = parse("# Раздел\n\n" + text, **kwargs).items
    return section.content


# ------------------------------------------------------------------------------- структура
def test_headings_fold_into_sections_and_nested_subsections():
    text = "# Цель работы\n\nа\n\n# Ход работы\n\n## Один\n\nб\n\n### Глубже\n\nв\n\n## Два\n\nг\n"
    first, second = parse(text).items
    assert first == Section("Цель работы", (P((Text("а"),)),))
    assert second.title == "Ход работы"
    one, two = second.content
    assert one == Subsection("Один", P((Text("б"),)), Subsection("Глубже", P((Text("в"),))))
    assert two == Subsection("Два", P((Text("г"),)))


def test_appendix_heading_uses_appendix_rule_and_kind():
    (section,) = parse("# Исходный код {.appendix status=справочное}\n\nтекст\n").items
    assert (section.rule_name, section.title, section.status) == ("Приложение", "Исходный код", "справочное")


def test_toc_and_pagebreak_divs():
    items = parse("::: {.toc levels=2}\n:::\n\n# Раздел\n\n::: pagebreak\n:::\n\nтекст\n").items
    assert items[0] == Toc(None, levels=2)
    assert items[1].content == (PageBreak(), P((Text("текст"),)))


def test_toc_between_sections_is_allowed_but_not_inside_subsection():
    items = parse("# А\n\nтекст\n\n::: toc\n:::\n\n# Б\n").items
    assert [type(item) for item in items] == [Section, Toc, Section]
    with pytest.raises(MarkdownError, match="Оглавление"):
        parse("# А\n\n## Под\n\n::: toc\n:::\n")


def test_text_before_first_heading_is_rejected():
    with pytest.raises(MarkdownError, match="до первого заголовка"):
        parse("просто текст\n")
    with pytest.raises(MarkdownError, match="первого уровня"):
        parse("## Подраздел\n")


def test_yaml_metadata_gives_title_and_author():
    document = parse("---\ntitle: ЛР4. *Регрессия*\nauthor: [А. Б. Иванов, В. Г. Петров]\n---\n\n# Раздел\n")
    assert (document.title, document.author) == ("ЛР4. Регрессия", "А. Б. Иванов, В. Г. Петров")


# ----------------------------------------------------------------------------------- инлайн
def test_inline_markup_matches_the_python_notation():
    (paragraph,) = blocks("Текст **жирный** и *курсив* и `код` и $x^2$.\n")
    assert paragraph == P(
        (
            Text("Текст "),
            B("жирный"),
            Text(" и "),
            I("курсив"),
            Text(" и "),
            Code("код"),
            Text(" и "),
            Math("x^2"),
            Text("."),
        )
    )


def test_nested_emphasis_collapses_into_flags():
    (paragraph,) = blocks("***оба*** и **жирный *с курсивом***\n")
    nodes = paragraph.parts[0]
    assert Text("оба", bold=True, italic=True) in nodes
    assert Text("с курсивом", bold=True, italic=True) in nodes


def test_superscript_subscript_hard_break_and_quotes():
    (paragraph,) = blocks("m^2^ и H~2~O\\\nновая \"строка\"\n")
    nodes = paragraph.parts[0]
    assert Text("2", script="sup") in nodes
    assert Text("2", script="sub") in nodes
    assert LineBreak() in nodes
    assert "".join(n.text for n in nodes if isinstance(n, Text)).endswith("«строка»")


def test_link_to_id_becomes_reference_and_external_link_keeps_address():
    (paragraph,) = blocks("См. [](#fig:x) и [сайт](https://example.com) и <https://a.b>.\n")
    nodes = paragraph.parts[0]
    assert Ref("fig:x") in nodes
    text = "".join(n.text for n in nodes if isinstance(n, Text))
    assert "сайт (https://example.com)" in text
    assert text.count("https://a.b") == 1


def test_inline_image_and_unsupported_inline_are_rejected():
    with pytest.raises(MarkdownError, match="отдельным абзацем"):
        blocks("текст ![п](x.png) текст\n")
    with pytest.raises(MarkdownError, match="Note"):
        blocks("текст[^1]\n\n[^1]: сноска\n")


# ------------------------------------------------------------------------------------- списки
def test_nested_lists_of_any_depth():
    (bullets,) = blocks("- один\n  - вложенный\n    - еще глубже\n- два\n")
    assert bullets == Bullets(
        (
            ListItem((Text("один"),), (ListItem((Text("вложенный"),), (ListItem((Text("еще глубже"),)),)),)),
            ListItem((Text("два"),)),
        )
    )


@pytest.mark.parametrize(
    ("source", "marker"),
    [("1) а\n2) б", "1)"), ("1. а\n2. б", "1."), ("a) а\nb) б", "a)"), ("a. а\nb. б", "a.")],
)
def test_numbered_list_markers(source, marker):
    (numbered,) = blocks(source + "\n")
    assert isinstance(numbered, Numbered) and numbered.marker == marker


def test_numbered_list_must_start_from_one_and_item_must_be_simple():
    with pytest.raises(MarkdownError, match="с 1"):
        blocks("3. а\n4. б\n")
    with pytest.raises(MarkdownError, match="один абзац"):
        blocks("- а\n\n  второй абзац\n")


# --------------------------------------------------------------------------- рисунки, листинги
def test_figure_with_label_width_and_path_relative_to_base_dir(tmp_path):
    (figure,) = blocks("![Подпись *рисунка*](img/x.png){#fig:x width=10cm}\n", base_dir=tmp_path)
    assert figure == Figure(
        str(tmp_path / "img" / "x.png"), (Text("Подпись "), I("рисунка")), label="fig:x", width_cm=10.0
    )


@pytest.mark.parametrize(("width", "cm"), [("100mm", 10.0), ("2in", 5.08), ("72pt", 2.54)])
def test_figure_width_units(width, cm, tmp_path):
    (figure,) = blocks(f"![п](x.png){{width={width}}}\n", base_dir=tmp_path)
    assert figure.width_cm == pytest.approx(cm)


def test_figure_width_in_percent_is_rejected(tmp_path):
    with pytest.raises(MarkdownError, match="cm, mm, in или pt"):
        blocks("![п](x.png){width=50%}\n", base_dir=tmp_path)


def test_listing_attributes():
    text = '``` {#lst:a .python .numberLines caption="Обучение" highlight=true}\nprint(1)\n```\n'
    (listing,) = blocks(text)
    assert listing == Listing(
        "print(1)", Text("Обучение"), label="lst:a", language="python", line_numbers=True, highlight=True
    )


def test_listing_without_attributes_follows_the_style():
    (listing,) = blocks("```\nx = 1\n```\n")
    assert listing == Listing("x = 1")
    (listing,) = blocks("```python\nx = 1\n```\n")
    assert (listing.language, listing.line_numbers, listing.highlight, listing.caption) == ("python", None, None, None)


def test_listing_can_be_read_from_a_file(tmp_path):
    (tmp_path / "code.py").write_text("a = 1\nb = 2\n", encoding="utf-8")
    (listing,) = blocks('``` {.python file=code.py caption="Из файла"}\n```\n', base_dir=tmp_path)
    assert listing.code.splitlines() == ["a = 1", "b = 2"]
    with pytest.raises(MarkdownError, match="пустым"):
        blocks("``` {file=code.py}\nтекст\n```\n", base_dir=tmp_path)


# --------------------------------------------------------------------------------------- формулы
def test_display_math_alone_is_a_formula_with_optional_label():
    plain, labeled = blocks("$$ a = b $$\n\n$$ \\frac{1}{n} $$ {#eq:mean}\n")
    assert plain == Formula("a = b")
    assert labeled == Formula("\\frac{1}{n}", label="eq:mean")


def test_where_block_attaches_definitions_to_the_formula():
    text = "$$ y = x $$ {#eq:y}\n\n::: where\n$y$\n:   фактическое *значение*\n\n$n$\n:   размер выборки\n:::\n"
    (formula,) = blocks(text)
    assert formula == replace(
        Formula("y = x", label="eq:y"),
        where=(("y", (Text("фактическое "), I("значение"))), ("n", (Text("размер выборки"),))),
    )


def test_where_block_without_formula_is_rejected():
    with pytest.raises(MarkdownError, match="после формулы"):
        blocks("текст\n\n::: where\n$y$\n:   описание\n:::\n")


def test_unknown_div_and_unsupported_block_give_clear_errors():
    with pytest.raises(MarkdownError, match="Неизвестный блок ::: note"):
        blocks("::: note\nтекст\n:::\n")
    with pytest.raises(MarkdownError, match="BlockQuote"):
        blocks("> цитата\n")


# --------------------------------------------------------------------------------------- таблицы
TABLE = "| Модель | MSE | R² |\n|:-------|----:|:--:|\n| Линейная | 3,2 | **0,81** |\n| Дерево | 2,9 | 0,84 |\n"


def test_bare_table_has_no_caption_and_keeps_column_alignment():
    (table,) = blocks(TABLE)
    assert table == Table(
        ((Text("Модель"),), (Text("MSE"),), (Text("R²"),)),
        [
            ((Text("Линейная"),), (Text("3,2"),), (B("0,81"),)),
            ((Text("Дерево"),), (Text("2,9"),), (Text("0,84"),)),
        ],
        align=("left", "right", "center"),
    )
    (table,) = blocks("| a | b |\n|---|---|\n| 1 | 2 |\n")
    assert table.align is None


def test_table_div_gives_caption_label_widths_and_font_size():
    text = f"::: {{.table #tbl:res widths=\"5,2,2\" font-size=12}}\nСравнение *моделей*\n\n{TABLE}:::\n"
    (table,) = blocks(text)
    assert table.caption == (Text("Сравнение "), I("моделей"))
    assert (table.label, table.col_widths_cm, table.font_size_pt) == ("tbl:res", (5.0, 2.0, 2.0), 12.0)
    after = blocks(f"::: {{.table #tbl:res}}\n{TABLE}\nПодпись после\n:::\n")[0]
    assert after.caption == (Text("Подпись после"),)


def test_table_problems_give_clear_errors():
    with pytest.raises(MarkdownError, match="одной таблицы"):
        blocks("::: table\nтолько текст\n:::\n")
    with pytest.raises(MarkdownError, match="список из 2 длин"):
        blocks("::: {.table widths=\"1,2,3,4,5\"}\n| a | b |\n|---|---|\n| 1 | 2 |\n:::\n")


def test_markdown_table_resolves_like_python_table_and_numbers_by_style():
    items = parse(f"# Раздел\n\n::: {{.table #tbl:res}}\nСравнение\n\n{TABLE}:::\n\nСм. [](#tbl:res).\n").items
    resolved = Report(SUAI, items=list(items)).resolve()
    table = next(node for node in resolved.nodes if isinstance(node, RTable))
    assert table.number is not None and table.number.display == "1"
    assert [cell.numeric for cell in table.rows[0]] == [False, True, True]
    assert table.align == ("left", "right", "center")
    assert resolved.labels["tbl:res"] == table.number


def test_markdown_table_builds_into_docx_with_bold_cell(tmp_path):
    path = tmp_path / "t.md"
    path.write_text(f"# Раздел\n\n::: {{.table #tbl:r}}\nПодпись\n\n{TABLE}:::\n", encoding="utf-8")
    from docx import Document

    doc = Document(str(Report.from_markdown(path, style=SUAI).build_to_docx(tmp_path / "t.docx")))
    (table,) = doc.tables
    assert [cell.text for cell in table.rows[1].cells] == ["Линейная", "3,2", "0,81"]
    assert table.rows[1].cells[2].paragraphs[0].runs[0].bold
    assert any(p.text.startswith("Таблица 1") for p in doc.paragraphs)


# ------------------------------------------------------------------- файлы и сквозная сборка
def test_report_from_files_resolves_the_same_as_python_notation(tmp_path):
    path = tmp_path / "report.md"
    path.write_text(
        "# Ход работы\n\nФормула [](#eq:a) и **важно**.\n\n$$ a=b $$ {#eq:a}\n\n"
        "- один\n- два\n\n``` {#lst:p .python caption=\"Код\"}\nx = 1\n```\n",
        encoding="utf-8",
    )
    from_markdown = Report.from_markdown(path, style=SUAI).resolve()
    by_hand = (
        Report(SUAI)
        .add_section(
            "Ход работы",
            P("Формула ", Ref("eq:a"), " и ", B("важно"), "."),
            Formula("a=b", label="eq:a"),
            Bullets(["один", "два"]),
            Listing("x = 1", Text("Код"), label="lst:p", language="python"),
        )
        .resolve()
    )
    assert from_markdown.nodes == by_hand.nodes
    assert from_markdown.labels == by_hand.labels


def test_report_from_several_files_keeps_order_and_takes_metadata_from_the_first(tmp_path):
    (tmp_path / "a.md").write_text("---\ntitle: Заголовок\n---\n\n# Введение\n", encoding="utf-8")
    (tmp_path / "b.md").write_text("# Заключение\n", encoding="utf-8")
    report = Report.from_markdown(tmp_path / "a.md", tmp_path / "b.md", style=SUAI, author="Автор")
    assert [item.title for item in report.items] == ["Введение", "Заключение"]
    assert (report.title, report.author) == ("Заголовок", "Автор")
    with pytest.raises(ValueError, match="хотя бы один"):
        Report.from_markdown(style=SUAI)


def test_markdown_report_builds_a_docx_with_figure_and_formula(tmp_path):
    (tmp_path / "x.png").write_bytes(png().getvalue())
    path = tmp_path / "r.md"
    path.write_text(
        "# Введение\n\nСм. [](#fig:x) и [](#eq:a).\n\n![Рисунок подпись](x.png){#fig:x}\n\n$$ a=b $$ {#eq:a}\n",
        encoding="utf-8",
    )
    out = Report.from_markdown(path, style=SUAI).build_to_docx(tmp_path / "out.docx")
    from docx import Document

    doc = Document(str(out))
    assert any("Рисунок 1" in p.text for p in doc.paragraphs)
    assert doc.element.body.findall(f".//{M}oMath")
    assert len(doc.inline_shapes) == 1


def test_parse_markdown_reads_file_relative_to_its_folder(tmp_path):
    (tmp_path / "x.png").write_bytes(png().getvalue())
    path = tmp_path / "r.md"
    path.write_text("# Раздел\n\n![](x.png)\n", encoding="utf-8")
    (section,) = parse_markdown(path).items
    assert section.content == (Figure(str(tmp_path / "x.png"), ()),)


def test_old_pandoc_api_version_is_rejected(monkeypatch):
    from studen_doc.frontends.markdown import pandoc_ast

    monkeypatch.setattr(pandoc_ast, "MIN_API_VERSION", (99, 0))
    with pytest.raises(MarkdownError, match="обновите pandoc"):
        parse("# Раздел\n")


# --------------------------------------------------------------------------------- метаданные
FRONT_MATTER = """---
style: suai
title_page:
  type: suai
  student_name: О. Н. Молодец
  group: "4318"
  year: 2026
  subject: ПРЕДМЕТ
  teacher_name: Иванов М. М.
  teacher_position: доцент
  work_number: 4
  title_bold: true
---

# Раздел
"""


def test_title_page_and_style_come_from_front_matter(tmp_path):
    path = tmp_path / "r.md"
    path.write_text(FRONT_MATTER, encoding="utf-8")
    report = Report.from_markdown(path)
    assert report.style is SUAI
    page = report.title_page
    assert isinstance(page, TitlePageSUAI)
    assert (page.student_name, page.group, page.year, page.work_number, page.title_bold) == (
        "О. Н. Молодец",
        "4318",
        2026,
        4,
        True,
    )
    assert page.teacher_position == "доцент"


def test_explicit_arguments_win_over_front_matter(tmp_path):
    path = tmp_path / "r.md"
    path.write_text(FRONT_MATTER, encoding="utf-8")
    own = TitlePageSUAI(student_name="Другой", group="1", subject="S", teacher_name="T", teacher_position="P")
    report = Report.from_markdown(path, style=GOST_732, title_page=own)
    assert report.style is GOST_732
    assert report.title_page is own


@pytest.mark.parametrize(
    ("meta", "message"),
    [
        ("title_page:\n  type: mgu\n", "известны"),
        ("title_page:\n  type: suai\n  nope: 1\n", "неизвестные поля"),
        ("title_page:\n  type: suai\n  student_name: X\n", "required"),
        ("title_page:\n  type: suai\n  year: abc\n  student_name: X\n", "не подходит"),
        ("title_page: suai\n", "словарем"),
    ],
)
def test_front_matter_errors_are_reported(meta, message, tmp_path):
    with pytest.raises(MarkdownError, match=message):
        parse(f"---\n{meta}---\n\n# Раздел\n")


def test_style_must_be_known_and_given_somewhere(tmp_path):
    path = tmp_path / "r.md"
    path.write_text("# Раздел\n", encoding="utf-8")
    with pytest.raises(MarkdownError, match="Стиль не задан"):
        Report.from_markdown(path)
    path.write_text("---\nstyle: nope\n---\n\n# Раздел\n", encoding="utf-8")
    with pytest.raises(MarkdownError, match="Неизвестный стиль"):
        Report.from_markdown(path)


# ------------------------------------------------------------------- смешивание с Python
def test_insert_places_python_blocks_and_lists_of_blocks():
    table = Table(["a"], [["1"]], caption="Из Python", label="tbl:py")
    text = "::: {.insert name=one}\n:::\n\nтекст\n\n::: {.insert name=many}\n:::\n"
    content = blocks(text, inserts={"one": table, "many": [P("раз"), P("два")]})
    assert content == (table, P((Text("текст"),)), P("раз"), P("два"))


def test_insert_unknown_name_is_reported():
    with pytest.raises(MarkdownError, match=r"'нет' не передан \(переданы: \['есть'\]\)"):
        blocks("::: {.insert name=нет}\n:::\n", inserts={"есть": P("x")})
    with pytest.raises(MarkdownError, match="нет имени блока"):
        blocks("::: insert\n:::\n")


def test_inserted_table_is_numbered_and_referenced_from_markdown():
    items = parse(
        "# Раздел\n\nСм. [](#tbl:py).\n\n::: {.insert name=t}\n:::\n",
        inserts={"t": Table(["a"], [["1"]], caption="Из Python", label="tbl:py")},
    ).items
    resolved = Report(SUAI, items=list(items)).resolve()
    assert resolved.labels["tbl:py"].display == "1"


def test_var_spans_are_formatted_with_decimal_comma_everywhere():
    text = (
        'R^2^ = []{var=r2 format=".3f"}, n = []{var=n}, имя []{var=name}, **[]{var=r2 format=".1f"}**.\n'
    )
    (paragraph,) = blocks(text, variables={"r2": 0.88123, "n": 506, "name": "лес"})
    shown = "".join(node.text for node in paragraph.parts[0] if isinstance(node, Text))
    assert shown == "R2 = 0,881, n = 506, имя лес, 0,9."
    (paragraph,) = blocks("x []{var=r2 format=\".2f\"}\n", variables={"r2": 0.5}, decimal=".")
    assert paragraph.parts[0] == (Text("x 0.50"),)


def test_var_span_with_unknown_name_lists_known_ones():
    with pytest.raises(MarkdownError, match=r"'q' не передана \(переданы: \['r2'\]\)"):
        blocks("[]{var=q}\n", variables={"r2": 1})


def test_report_from_markdown_passes_blocks_and_values(tmp_path):
    path = tmp_path / "r.md"
    path.write_text(
        '# Раздел\n\nЛучшая модель: []{var=best}, R^2^ = []{var=r2 format=".2f"}.\n\n::: {.insert name=tbl}\n:::\n',
        encoding="utf-8",
    )
    report = Report.from_markdown(
        path,
        style=SUAI,
        variables={"best": "лес", "r2": 0.873},
        inserts={"tbl": Table(["a"], [["1"]], caption="Итог", label="tbl:r")},
    )
    resolved = report.resolve()
    assert resolved.labels["tbl:r"].display == "1"
    paragraph = next(node for node in resolved.nodes if type(node).__name__ == "RParagraph")
    assert "".join(n.text for n in paragraph.inlines if isinstance(n, Text)).endswith("R2 = 0,87.")


def test_keep_with_next_div_marks_the_paragraphs_inside():
    text = "::: keep-with-next\n**Вопрос 1**\n:::\n\nОтвет.\n"
    question, answer = blocks(text)
    assert question.keep_with_next is True and answer.keep_with_next is False
    assert question == P((B("Вопрос 1"),), keep_with_next=True)  # то же содержимое, что у абзаца без обертки


def test_substituted_negative_numbers_get_a_minus_sign_unless_asked_otherwise():
    text = 'R2 = []{var=r2 format=".3f"}\n'
    (default,) = blocks(text, variables={"r2": -0.3461})
    (plain,) = blocks(text, variables={"r2": -0.3461}, minus="-")
    assert default.parts[0][0].text == "R2 = −0,346"
    assert plain.parts[0][0].text == "R2 = -0,346"
